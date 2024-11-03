from pathlib import Path
from docxtpl import DocxTemplate
import openpyxl
import pandas as pd
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx import Document
from docx.shared import Cm
import os
import numpy as np
from datetime import date
import datetime as dt
from PyInquirer import prompt
import pprint
import tkinter as tk


base_dir = current_directory = os.getcwd()
parent_dir = os.path.dirname(base_dir)
supparentdir = os.path.dirname(parent_dir)
template_path = os.path.join(parent_dir,"Vorlage.docx")
excel_template_path = os.path.join(parent_dir,"Jahresübersicht_Vorlage.xlsx")
allhourdata_path = os.path.join(parent_dir ,"Stundendaten.xlsx")
allclientdata_path = os.path.join(parent_dir ,"PatientInneninformationen.xlsx")

outputdir_path = os.path.join(supparentdir,f"{dt.datetime.now().year}")

if not os.path.isdir(outputdir_path):
    os.mkdir(outputdir_path)
else:
    print(f"We already have a directory {outputdir_path}")


outputfile_path = os.path.join(outputdir_path, f"RE {1} {dt.date.today().strftime('%d_%m_%Y')}.docx")

# invoicedata has to be a dict with keys which are the same as the placeholders in the template
# input the client data  in word
doc = DocxTemplate(template_path)
doc.render(invoicedata)
doc.save(outputfile_path)

## input the hour table in word
doc = Document(outputfile_path)
doc.tables #a list of all tables in document
# table nr. 0 is the data table and table nr. 1 is the sum table

# change a table in the template
wordtable = pd.concat([namehourdata["Datum"].apply(lambda x: x.strftime("%d.%m.%Y")), namehourdata["Minuten"].apply(lambda x: str(x) + " min"), amountpersession.apply(lambda x: "%0.2f" % x + " €") ], axis=1)
print("Die Stunden sind: " )
print(wordtable)


# insert the table in the Word document
for index, row in wordtable.iterrows():
    hourdatatable = doc.tables[0]   #so hourdatatable is the first table in the document
    data_row = hourdatatable.add_row().cells
    for i,(name,entry) in enumerate(row.items()):
            data_row[i].text = entry
#format it
for row in doc.tables[0].rows:
    row.height = Cm(0.8)
    row.alignment = WD_TABLE_ALIGNMENT.CENTER



#insert total amount into tables[1]
totalamount = sum(np.array(amountpersession))
totalamountstring = (str(totalamount)+"0").replace(".",",")
doc.tables[1].cell(0, 2).text = str(totalamount) + "0" + " €"

doc.save(outputfile_path)


os.startfile(outputfile_path)